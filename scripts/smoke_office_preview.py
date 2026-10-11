"""Opt-in Windows/Office smoke: shared PowerPoint, cancellation and cleanup.

Creates and closes only a hidden sentinel presentation; existing user documents
and the existing PowerPoint process must survive every path.
"""
import base64
import json
import os
from pathlib import Path
import queue
import subprocess
import sys
import tempfile
import threading
import time
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from core import document_preview
from core.process_utils import subprocess_kwargs_no_window
from core.workbench import WorkbenchError


HOLDER = r'''
$ErrorActionPreference = 'Stop'
$ProgressPreference = 'SilentlyContinue'
[Console]::OutputEncoding = New-Object System.Text.UTF8Encoding($false)
$existed = @(Get-Process POWERPNT -ErrorAction SilentlyContinue).Count -gt 0
$app = $null; $guard = $null
try {
    $app = [Activator]::CreateInstance([Type]::GetTypeFromProgID('PowerPoint.Application'))
    $guard = $app.Presentations.Add(0)
    $slide = $guard.Slides.Add(1, 1)
    $slide.Shapes.Title.TextFrame.TextRange.Text = 'Unsaved preview guard'
    while ($true) {
        [Console]::Out.WriteLine((@{count=$app.Presentations.Count; saved=[int]$guard.Saved;
            text=$guard.Slides.Item(1).Shapes.Title.TextFrame.TextRange.Text;
            security=[int]$app.AutomationSecurity; alerts=[int]$app.DisplayAlerts} | ConvertTo-Json -Compress))
        [Console]::Out.Flush()
        if ([Console]::In.ReadLine() -ne 'state') { break }
    }
} catch {
    [Console]::Out.WriteLine('Holder error: ' + $_.Exception.Message + ' / ' + $_.InvocationInfo.PositionMessage)
    exit 1
} finally {
    if ($null -ne $guard) { $guard.Saved = -1; $guard.Close() }
    if ($null -ne $app) {
        if (-not $existed -and $app.Presentations.Count -eq 0) { $app.Quit() }
        [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($app)
    }
}
'''


def wait(predicate, timeout=15):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if predicate(): return
        time.sleep(.05)
    raise AssertionError("Office cleanup did not finish")


def run():
    from pptx import Presentation
    from pypdf import PdfReader
    powershell = str(Path(os.environ["SystemRoot"]) / "System32/WindowsPowerShell/v1.0/powershell.exe")
    holder = subprocess.Popen([powershell, "-NoLogo", "-NoProfile", "-NonInteractive", "-STA", "-EncodedCommand",
                               base64.b64encode(HOLDER.encode("utf-16-le")).decode("ascii")],
                              stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                              **subprocess_kwargs_no_window())
    states = queue.Queue()
    def read():
        for line in iter(holder.stdout.readline, b""):
            states.put(line.decode("utf-8", errors="replace").strip())
    threading.Thread(target=read, daemon=True).start()
    def state(initial=False):
        if not initial:
            holder.stdin.write(b"state\n"); holder.stdin.flush()
        message = states.get(timeout=15)
        try:
            return json.loads(message)
        except ValueError as exc:
            raise AssertionError(f"PowerPoint test holder failed: {message}") from exc
    try:
        baseline = state(initial=True)
        assert baseline["saved"] == 0 and baseline["text"] == "Unsaved preview guard"
        with tempfile.TemporaryDirectory(prefix="cowork-office-smoke-") as temporary:
            root = Path(temporary)
            source = root / "preview.pptx"
            deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[0])
            slide.shapes.title.text = "Office preview regression"
            deck.save(source)
            original = source.read_bytes()
            cache = root / "cache"; cache.mkdir()
            pdf, _ = document_preview.render_office_pdf(source, root=cache, force=True)
            assert Path(pdf).read_bytes().startswith(b"%PDF-")
            pages = PdfReader(pdf).pages
            assert len(pages) == 1 and "Office preview regression" in pages[0].extract_text()
            assert source.read_bytes() == original and state() == baseline
            print("PASS shared PowerPoint render preserves unsaved presentation and application settings", flush=True)
            Path(pdf).unlink()
            broken = root / "broken.pptx"; broken.write_bytes(b"invalid presentation")
            try:
                document_preview.render_office_pdf(broken, root=cache)
                raise AssertionError("invalid presentation was accepted")
            except WorkbenchError as exc:
                assert exc.code == "renderer_failed", exc
            assert state() == baseline and not list(cache.iterdir())
            print("PASS failed render preserves shared PowerPoint and cleans temporary files", flush=True)

            for timing in ("early", "opened", "timeout"):
                cancel = threading.Event()
                marker = root / "opened"
                script = document_preview._OFFICE_SCRIPT
                if timing != "early":
                    # Hold a real, already-opened document at a deterministic
                    # boundary to exercise cancellation while COM is occupied.
                    script = script.replace("        Assert-Active\n        $doc.SaveAs",
                        "        [IO.File]::WriteAllText('" + str(marker).replace("'", "''") + "', 'opened')\n"
                        "        Start-Sleep -Seconds 2\n        Assert-Active\n        $doc.SaveAs")
                def cancel_after_open():
                    wait(marker.exists)
                    cancel.set()
                trigger = threading.Timer(.1, cancel.set) if timing == "early" else threading.Thread(target=cancel_after_open)
                if timing != "timeout": trigger.start()
                try:
                    with patch.object(document_preview, "_OFFICE_SCRIPT", script):
                        document_preview.render_office_pdf(source, root=cache, force=True,
                                                           cancel=cancel, timeout=1.5 if timing == "timeout" else 30)
                    raise AssertionError("cancelled conversion returned success")
                except WorkbenchError as exc:
                    assert exc.code == ("renderer_timeout" if timing == "timeout" else "cancelled"), exc
                finally:
                    if timing != "timeout": trigger.join(15)
                wait(lambda: not list(cache.iterdir()))
                assert state() == baseline and source.read_bytes() == original
                marker.unlink(missing_ok=True)
                print(f"PASS {timing} cancellation/timeout closes only the preview copy", flush=True)
    finally:
        if holder.poll() is None:
            holder.stdin.write(b"close\n"); holder.stdin.flush()
        holder.wait(timeout=15)
        holder.stdin.close(); holder.stdout.close()


if __name__ == "__main__":
    run()
