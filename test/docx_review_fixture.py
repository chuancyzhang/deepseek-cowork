"""A standards-shaped DOCX with real note relationships and tracked text."""
import io
import zipfile
from lxml import etree
from docx import Document
from core.docx_review import W, tag


def make_review_docx(path):
    document = Document()
    document.add_heading("文档审阅示例", 0)
    paragraph = document.add_paragraph("正文原文，脚注")
    reference = etree.SubElement(paragraph.add_run()._r, tag("footnoteReference"))
    reference.set(tag("id"), "1")
    paragraph.add_run("，尾注")
    reference = etree.SubElement(paragraph.add_run()._r, tag("endnoteReference"))
    reference.set(tag("id"), "2")
    paragraph.add_run("，普通内容。")
    paragraph = document.add_paragraph("修订前后：")
    for name, identifier, text in (("ins", "10", "新增文字"), ("del", "11", "删除文字")):
        change = etree.SubElement(paragraph._p, tag(name), {tag("id"): identifier, tag("author"): "审阅人", tag("date"): "2026-10-13T08:00:00Z"})
        run = etree.SubElement(change, tag("r"))
        etree.SubElement(run, tag("t" if name == "ins" else "delText")).text = text
    paragraph.add_run("。")
    paragraph = document.add_paragraph()
    paragraph.add_run("保留粗体").bold = True
    paragraph.add_run("和普通文字。")
    paragraph = document.add_paragraph("格式修订段落")
    properties = paragraph._p.get_or_add_pPr()
    etree.SubElement(properties, tag("jc"), {tag("val"): "right"})
    change = etree.SubElement(properties, tag("pPrChange"), {tag("id"): "12", tag("author"): "审阅人"})
    previous = etree.SubElement(change, tag("pPr"))
    etree.SubElement(previous, tag("jc"), {tag("val"): "center"})
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "表格原文"
    table.cell(0, 1).text = "第二列"
    document.sections[0].header.paragraphs[0].text = "保留的页眉"
    buffer = io.BytesIO()
    document.save(buffer)
    with zipfile.ZipFile(buffer) as archive:
        parts = {info.filename: archive.read(info) for info in archive.infolist()}
    relation_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    content_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    relationships = etree.fromstring(parts["word/_rels/document.xml.rels"])
    content_types = etree.fromstring(parts["[Content_Types].xml"])
    for kind, identifier, text in (("footnote", "1", "脚注原文。"), ("endnote", "2", "尾注原文。")):
        root = etree.Element(tag(kind + "s"), nsmap={"w": W})
        for note_id, note_type in (("-1", "separator"), ("0", "continuationSeparator")):
            note = etree.SubElement(root, tag(kind), {tag("id"): note_id, tag("type"): note_type})
            etree.SubElement(etree.SubElement(etree.SubElement(note, tag("p")), tag("r")), tag(note_type))
        note = etree.SubElement(root, tag(kind), {tag("id"): identifier})
        paragraph = etree.SubElement(note, tag("p"))
        etree.SubElement(etree.SubElement(paragraph, tag("r")), tag(kind + "Ref"))
        etree.SubElement(etree.SubElement(paragraph, tag("r")), tag("t")).text = text
        if kind == "footnote":
            change = etree.SubElement(paragraph, tag("ins"), {tag("id"): "13", tag("author"): "脚注审阅人"})
            etree.SubElement(etree.SubElement(change, tag("r")), tag("t")).text = "脚注新增文字"
        parts[f"word/{kind}s.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
        etree.SubElement(relationships, f"{{{relation_ns}}}Relationship", {
            "Id": f"rId{kind}", "Type": f"http://schemas.openxmlformats.org/officeDocument/2006/relationships/{kind}s", "Target": f"{kind}s.xml"})
        etree.SubElement(content_types, f"{{{content_ns}}}Override", {
            "PartName": f"/word/{kind}s.xml", "ContentType": f"application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}s+xml"})
    parts["word/_rels/document.xml.rels"] = etree.tostring(relationships)
    parts["[Content_Types].xml"] = etree.tostring(content_types)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return parts
